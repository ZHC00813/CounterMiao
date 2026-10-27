import os, sys, re, json, hashlib, shutil, subprocess, tempfile, threading, queue, zipfile, csv
from pathlib import Path
from decimal import Decimal, InvalidOperation
from datetime import datetime
from copy import deepcopy
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from tkinterdnd2 import TkinterDnD, DND_FILES, COPY
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from PIL import Image, ImageOps
from io import BytesIO
import pymupdf

BASE=Path(getattr(sys,'_MEIPASS',Path(__file__).parent))
SUPPORTED={'.pdf','.png','.jpg','.jpeg','.bmp','.tif','.tiff'}
def collect_invoices(inputs):
    files=[]; seen=set(); skipped=[]
    def add(path):
        try:
            if path.is_file() and path.suffix.lower() in SUPPORTED:
                resolved=str(path.resolve());key=os.path.normcase(resolved)
                if key not in seen:seen.add(key);files.append(resolved)
            else:skipped.append(str(path))
        except OSError:skipped.append(str(path))
    for raw in inputs:
        path=Path(raw)
        try:
            if path.is_dir():
                for parent,dirs,names in os.walk(path,followlinks=False,onerror=lambda e:skipped.append(str(e.filename))):
                    dirs[:]=sorted(d for d in dirs if not Path(parent,d).is_symlink())
                    for name in sorted(names):add(Path(parent,name))
            else:add(path)
        except OSError:skipped.append(str(path))
    return files,skipped

def amount_summary(invoices):
    amount=Decimal(0);known=0;alerts=0
    for inv in invoices:
        try:value=money(inv.get('amount',''))
        except (ValueError,InvalidOperation):continue
        amount+=value;known+=1;alerts+=value>1000
    return dict(amount=amount,known=known,unknown=len(invoices)-known,count=len(invoices),alerts=alerts)

def write_amount_list(invoices,path):
    stats=amount_summary(invoices)
    with open(path,'w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f);writer.writerow(['序号','文件名','发票号码','含税金额（元）','金额状态','超过1000元'])
        for n,inv in enumerate(invoices,1):
            try:value=money(inv.get('amount',''));amount=f'{value:.2f}';state='已核对' if inv.get('confirmed') else '待核对';alert='是' if value>1000 else ''
            except (ValueError,InvalidOperation):amount='';state='金额待识别 / 补录';alert=''
            writer.writerow([n,Path(inv['source']).name,inv.get('number',''),amount,state,alert])
        writer.writerow(['已识别金额合计','','',f'{stats["amount"]:.2f}',f'另有{stats["unknown"]}张金额待识别',''])
def money(x):
    d=Decimal(str(x).replace(',','').replace('￥','').replace('¥','').strip())
    if not d.is_finite() or d<0 or d!=d.quantize(Decimal('.01')): raise ValueError('金额须为非负数，最多两位小数')
    return d
def total(inv): return sum((money(i['amount']) for i in inv['items']),Decimal(0))
def validate(inv):
    if not inv.get('number'): raise ValueError('请填写发票号码')
    if inv.get('buyer')!='深圳科创学院': raise ValueError('请核对购买方抬头，须为深圳科创学院')
    if not inv['items']: raise ValueError('请填写真实采购项目')
    for i in inv['items']:
        if not i['name'].strip() or not i['qty'].strip(): raise ValueError('品名、数量不能为空')
        if Decimal(i['qty'])<=0: raise ValueError('数量必须大于零')
        money(i['amount'])
    if any(not Path(p).is_file() for p in inv.get('photos',[])): raise ValueError('图片不存在，请重新关联或清除')
    if total(inv)!=money(inv['amount']): raise ValueError('逐项含税金额之和与整票金额不一致')
    if not Path(inv['source']).is_file(): raise ValueError('原发票文件不存在')
    if not inv.get('confirmed'): raise ValueError('请核对并确认识别结果')
def group(invoices):
    bins=[]; blocked=[]
    def key(x):
        try:return money(x['amount'])
        except Exception:return Decimal(-1)
    for inv in sorted(invoices,key=key,reverse=True):
        try:
            validate(inv)
            if money(inv['amount'])>5000 or len(inv['items'])>15: raise ValueError('单票超过5000元或15项，需单独处理，不自动拆票')
        except Exception as e: blocked.append((inv,str(e))); continue
        for b in bins:
            if b['amount']+money(inv['amount'])<=5000 and b['count']+len(inv['items'])<=15 and b['category']==inv.get('category','采购'):
                b['invoices'].append(inv); b['amount']+=money(inv['amount']); b['count']+=len(inv['items']); break
        else: bins.append(dict(invoices=[inv],amount=money(inv['amount']),count=len(inv['items']),category=inv.get('category','采购')))
    return bins,blocked
def ocr(path):
    r=subprocess.run(['powershell.exe','-NoProfile','-ExecutionPolicy','Bypass','-File',str(BASE/'ocr.ps1'),str(path)],capture_output=True,creationflags=0x08000000,timeout=90)
    if r.returncode: raise ValueError('Windows OCR 未成功，请检查中文语言包或手动核对录入')
    return r.stdout.decode('utf-8-sig',errors='replace')
def extract(path):
    path=Path(path); texts=[]; structured=[]; previous_item=None
    if path.suffix.lower()=='.pdf':
        with pymupdf.open(path) as pdf:
            for page in pdf:
                t=page.get_text(sort=True)
                if len(t.strip())<40:
                    with tempfile.TemporaryDirectory() as tmp:
                        p=Path(tmp)/'page.png'; page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5)).save(p); t=ocr(p)
                texts.append(t)
                # Preserve geometric rows for common digital VAT invoice layouts.
                rows=[]
                for word in sorted(page.get_text('words'),key=lambda w:(round(w[1]/3),w[0])):
                    target=next((r for r in rows if abs(r[0]-word[1])<3),None)
                    if target is None:target=[word[1],[]];rows.append(target)
                    target[1].append(word)
                for _,words in rows:
                    line=' '.join(w[4] for w in sorted(words,key=lambda w:w[0]))
                    m=re.fullmatch(r'\s*(\*.+?)\s+([^\s]+)\s+(\d+(?:\.\d+)?)\s+(\d+(?:\.\d+)?)\s+([\d,]+\.\d{2})\s+(\d+%|免税|不征税|零税率)\s*([\d,]+\.\d{2}|\*+)?\s*',line)
                    if m:
                        tax=m.group(7);tax=Decimal(tax.replace(',','')) if tax and '*' not in tax else Decimal(0)
                        previous_item=dict(name=m.group(1),model='',qty=m.group(3),amount=str(Decimal(m.group(5).replace(',',''))+tax),photos=[])
                        structured.append(previous_item)
                        continue
                    m=re.fullmatch(r'\s*(\*.+?)\s+(\d+(?:\.\d+)?)\s+(\d+(?:\.\d+)?)\s+([\d,]+\.\d{2})\s+(\d+%|免税|不征税|零税率)\s*([\d,]+\.\d{2}|\*+)?\s*',line)
                    if m:
                        tax=m.group(6);tax=Decimal(tax.replace(',','')) if tax and '*' not in tax else Decimal(0)
                        previous_item=dict(name=m.group(1),model='',qty=m.group(2),amount=str(Decimal(m.group(4).replace(',',''))+tax),photos=[])
                        structured.append(previous_item)
                    elif line.lstrip().startswith('*'):
                        discount=re.fullmatch(r'\s*(\*.+?)\s+(-[\d,]+\.\d{2})\s+(\d+%|免税|不征税|零税率)\s*(-[\d,]+\.\d{2}|\*+)?\s*',line)
                        if discount and previous_item and re.sub(r'\s+','',previous_item['name']).startswith(re.sub(r'\s+','',discount.group(1))):
                            tax=discount.group(4);tax=Decimal(tax.replace(',','')) if tax and '*' not in tax else Decimal(0)
                            adjusted=Decimal(previous_item['amount'])+Decimal(discount.group(2).replace(',',''))+tax
                            if adjusted>=0:previous_item['amount']=str(adjusted)
                        previous_item=None
    else:
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'image.png'; im=Image.open(path); im.thumbnail((2400,2400)); im.convert('RGB').save(p); texts=[ocr(p)]
    t='\n'.join(texts); flat=re.sub(r'[ \t]+','',t)
    number=re.search(r'(?:发票号码|发票号)[:：\s]*(\d{8,20})',flat)
    amounts=re.findall(r'[¥￥]\s*([\d,]+\.\d{2})',t)
    amount=amounts[-1].replace(',','') if amounts else ''
    m=re.search(r'(?:小写|价税合计)[）)\s:：]*[¥￥]?\s*([\d,]+\.\d{2})',flat)
    if m: amount=m.group(1).replace(',','')
    names=[]
    for line in t.splitlines():
        if '*' in line:
            name=re.sub(r'\s+',' ',line).strip()
            names.append(dict(name=name,model='',qty='',amount='',photos=[]))
    return dict(source=str(path.resolve()),hash=hashlib.sha256(path.read_bytes()).hexdigest(),number=number.group(1) if number else '',buyer='深圳科创学院' if '深圳科创学院' in flat else '',amount=amount,category='采购',items=structured or names,text=t,confirmed=False)

def unique_cells(row):
    out=[]
    for c in row.cells:
        if not out or c._tc is not out[-1]._tc: out.append(c)
    return out
def build_doc(batch,path):
    d=Document(BASE/'template.docx'); old=d.tables[0]
    table=d.add_table(rows=9,cols=7);table.autofit=False
    from docx.oxml.ns import qn
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');borders.append(el)
    table._tbl.tblPr.append(borders)
    widths=[.38,1.18,.92,.55,.65,.75,1.17]
    for row in table.rows:
        for c,w in zip(row.cells,widths): c.width=Inches(w)
    headers=['序号','货物/服务名称','型号参数','送达数量','单价（含税）','金额小计','其他（发票号）']
    for c,v in zip(table.rows[0].cells,headers):c.text=v
    table.rows[2].cells[0].merge(table.rows[2].cells[1]).text='金额合计'
    table.rows[2].cells[2].merge(table.rows[2].cells[6])
    table.rows[3].cells[0].merge(table.rows[3].cells[6]).text='验收照片'
    table.rows[4].cells[0].merge(table.rows[4].cells[6])
    for r,label in [(5,'验收内容'),(6,'验收时间'),(7,'团队经办人\n签字确认'),(8,'辅导老师\n签字确认')]:
        table.rows[r].cells[0].merge(table.rows[r].cells[1]).text=label
        table.rows[r].cells[2].merge(table.rows[r].cells[6]).text='年    月    日' if r==6 else ''
    old._tbl.addprevious(table._tbl);old._tbl.getparent().remove(old._tbl)
    sample=deepcopy(table.rows[1]._tr);table._tbl.remove(table.rows[1]._tr)
    items=[(inv,i) for inv in batch['invoices'] for i in inv['items']]
    for n,(inv,item) in enumerate(items,1):
        row=deepcopy(sample)
        for el in row.iter():
            for a in list(el.attrib):
                if a.endswith('}paraId') or a.endswith('}textId'): del el.attrib[a]
        table.rows[n]._tr.addprevious(row)
        cells=unique_cells(table.rows[n]); qty=Decimal(item['qty']); unit=money(item['amount'])/qty
        values=[str(n),item['name'],item['model'],item['qty'],f'{unit:.4f}'.rstrip('0').rstrip('.'),f'{money(item["amount"]):.2f}',inv['number']]
        for c,v in zip(cells,values): c.text=v
        pr=row.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
    unique_cells(table.rows[1+len(items)])[1].text=f'RMB {batch["amount"]:.2f}（含税）'
    photo_row=table.rows[3+len(items)]; photo_sample=deepcopy(photo_row._tr)
    conclusion=unique_cells(table.rows[4+len(items)])[1]
    # Compact rows contain only real photos; the final row has no empty slots.
    photos=[(inv,photo) for inv in batch['invoices'] for photo in inv.get('photos',[])]
    photo_count=len(photos)
    from docx.table import _Row
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    for start in range(0,photo_count,4):
        chunk=photos[start:start+4];row=deepcopy(photo_sample)
        for el in row.iter():
            for a in list(el.attrib):
                if a.endswith('}paraId') or a.endswith('}textId'):del el.attrib[a]
        photo_row._tr.addprevious(row)
        cell=unique_cells(_Row(row,table))[0];cell.text=''
        pr=row.get_or_add_trPr()
        for h in list(pr):
            if h.tag.endswith('}trHeight'):pr.remove(h)
        pr.append(OxmlElement('w:cantSplit'))
        grid=cell.add_table(rows=1,cols=len(chunk));grid.autofit=False
        for col in grid.columns:col.width=Inches(1.38)
        grid.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for tile,(inv,photo) in zip(grid.rows[0].cells,chunk):
            tile.width=Inches(1.38)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                m=OxmlElement('w:'+edge);m.set(qn('w:w'),'30');m.set(qn('w:type'),'dxa');margins.append(m)
            tile._tc.get_or_add_tcPr().append(margins)
            p=tile.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(2);p.paragraph_format.keep_with_next=True
            with Image.open(photo) as original:
                im=ImageOps.exif_transpose(original).convert('RGBA' if 'A' in original.getbands() else 'RGB')
                w,h=im.size;scale=min(1.28/w,1.25/h)
                picture=BytesIO();im.save(picture,format='PNG');picture.seek(0)
                p.add_run().add_picture(picture,width=Inches(w*scale),height=Inches(h*scale))
            p=tile.add_paragraph(f'发票 {inv["number"]}\n'+'、'.join(i['name'] for i in inv['items']))
            p.paragraph_format.keep_together=True;p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
            for run in p.runs:run.font.size=Pt(8)
        for p in cell.paragraphs:
            p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(1)
            for run in p.runs:run.font.size=Pt(1)
    table._tbl.remove(photo_row._tr)
    if not photo_count:
        title_row=table.rows[2+len(items)];table._tbl.remove(title_row._tr)
    else:
        for p in unique_cells(table.rows[2+len(items)])[0].paragraphs:
            p.paragraph_format.keep_with_next=True
    conclusion.text='□ 合格    □ 不合格'
    for row_index,row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in unique_cells(row):
            for p in cell.paragraphs:
                if len(table.rows)-4<=row_index<len(table.rows)-1:p.paragraph_format.keep_with_next=True
                for r in p.runs: r.font.size=Pt(10)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    # Word requires unique nonzero paragraph identifiers after template cloning.
    for idx,el in enumerate(d._element.iter(),1):
        for a in list(el.attrib):
            if a.endswith('}paraId') or a.endswith('}textId'): el.set(a,f'{idx:08X}')
    # The retained template also contains a header drawing. Word requires drawing
    # identifiers to be unique across document parts, including that header.
    drawing_id=1
    for part in d.part.package.parts:
        element=getattr(part,'_element',None)
        if element is not None:
            for drawing in element.iter(qn('wp:docPr')):
                drawing.set('id',str(drawing_id));drawing_id+=1
    d.save(path)
def export(batches,root):
    dest=Path(root)/('采购验收_'+datetime.now().strftime('%Y%m%d_%H%M%S_%f')); dest.mkdir()
    summary=[]
    for n,b in enumerate(batches,1):
        folder=dest/f'验收单{n:02}'; folder.mkdir(); build_doc(b,folder/'采购验收单.docx')
        mapping=[]
        for inv in b['invoices']:
            src=folder/'发票'; src.mkdir(exist_ok=True); shutil.copy2(inv['source'],src/(inv['hash'][:10]+'_'+Path(inv['source']).name))
            for k,i in enumerate(inv['items'],1):
                pics=[]
                for j,p in enumerate(inv.get('photos',[]),1):
                    ph=folder/'照片'; ph.mkdir(exist_ok=True); target=ph/f'{inv["hash"][:10]}_{k}_{j}{Path(p).suffix}'; shutil.copy2(p,target); pics.append(str(target.relative_to(folder)))
                mapping.append(dict(发票号=inv['number'],品名=i['name'],含税金额=i['amount'],照片=pics))
        (folder/'对应清单.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2),encoding='utf-8')
        summary.append(dict(验收单=n,项目数=b['count'],含税金额=str(b['amount']),发票=[v['number'] for v in b['invoices']],版式='待用户在Word中预览核对，未自动视觉核验'))
    (dest/'总清单.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    shutil.make_archive(str(dest),'zip',dest)
    return dest

class App(TkinterDnD.Tk):
    def __init__(self):
        super().__init__(); self.title('采购验收助手 · 2.0'); self.geometry('1240x810'); self.minsize(1100,720); self.invoices=[]; self.q=queue.Queue(); self.project=None
        self.importing=False;self.pending_paths=set();self.stats={}
        self.configure(background='white');self.option_add('*Font',('Microsoft YaHei UI',10))
        style=ttk.Style(); style.theme_use('clam')
        style.configure('.',background='white',foreground='#202824',font=('Microsoft YaHei UI',10))
        style.configure('TFrame',background='white');style.configure('TLabel',background='white')
        style.configure('TButton',padding=(12,9),background='#217346',foreground='white',bordercolor='#185c37',focuscolor='#185c37')
        style.map('TButton',background=[('active','#185c37'),('pressed','#10452a'),('disabled','#6d766f')],foreground=[('disabled','white')])
        style.configure('Treeview',rowheight=36,font=('Microsoft YaHei UI',10),background='white',fieldbackground='white',foreground='#202824',bordercolor='#cbd1cd')
        style.configure('Treeview.Heading',background='#f1f3f2',foreground='#263d30',font=('Microsoft YaHei UI',10,'bold'),padding=9)
        style.map('Treeview',background=[('selected','#217346')],foreground=[('selected','white')])
        header=tk.Frame(self,bg='#185c37');header.pack(fill='x')
        tk.Label(header,text='采购验收助手',font=('Microsoft YaHei UI',22,'bold'),bg='#185c37',fg='white').pack(side='left',padx=24,pady=17)
        tk.Label(header,text='发票整理  /  金额清单  /  采购验收',bg='#185c37',fg='white').pack(side='right',padx=24)
        bar=ttk.Frame(self); bar.pack(fill='x',padx=24,pady=(18,12))
        for text,cmd in [('导入发票',self.add),('导入文件夹',self.add_folder),('核对 / 配图',self.edit),('查看原票',self.open_source),('移除',self.remove),('保存项目',self.save),('打开项目',self.load),('金额清单',self.save_amounts),('分组导出',self.generate)]: ttk.Button(bar,text=text,command=cmd).pack(side='left',padx=(0,6))
        self.drop_zone=tk.Frame(self,bg='white',highlightbackground='#217346',highlightthickness=2);self.drop_zone.pack(fill='x',padx=24,pady=(0,14))
        tk.Label(self.drop_zone,text='把发票文件或整个文件夹拖到这里，即可导入',font=('Microsoft YaHei UI',14,'bold'),fg='#217346',bg='white').pack(pady=(16,4))
        tk.Label(self.drop_zone,text='支持多选拖入 · 自动读取子文件夹 · 也可拖到下方发票列表或窗口空白处',bg='white',fg='#59615c').pack(pady=(0,16))
        summary=ttk.Frame(self);summary.pack(fill='x',padx=24,pady=(0,10))
        self.total_text=tk.StringVar(value='已识别金额合计  ¥ 0.00')
        ttk.Label(summary,textvariable=self.total_text,font=('Microsoft YaHei UI',20,'bold'),foreground='#217346').pack(side='left')
        self.count_text=tk.StringVar(value='共 0 张发票');ttk.Label(summary,textvariable=self.count_text).pack(side='right')
        ttk.Label(self,text='每份 ≤ 5000 元 · 最多 15 个采购项目 · 超过 1000 元橙色提醒 · 无图不留空位',foreground='#666d68').pack(anchor='w',padx=24)
        listing=ttk.Frame(self);listing.pack(fill='both',expand=True,padx=24,pady=12)
        self.tree=ttk.Treeview(listing,columns=('file','number','amount','count','photos','status'),show='headings',selectmode='browse')
        for c,t,w in [('file','发票文件',210),('number','发票号码',205),('amount','含税金额（元）',140),('count','项目数',65),('photos','图片',70),('status','待办与提醒',360)]:
            self.tree.heading(c,text=t);self.tree.column(c,width=w,minwidth=60,anchor='e' if c=='amount' else 'w')
        ys=ttk.Scrollbar(listing,orient='vertical',command=self.tree.yview);xs=ttk.Scrollbar(listing,orient='horizontal',command=self.tree.xview)
        self.tree.configure(yscrollcommand=ys.set,xscrollcommand=xs.set)
        self.tree.grid(row=0,column=0,sticky='nsew');ys.grid(row=0,column=1,sticky='ns');xs.grid(row=1,column=0,sticky='ew');listing.rowconfigure(0,weight=1);listing.columnconfigure(0,weight=1)
        self.tree.tag_configure('warn',background='#fff0ce'); self.tree.bind('<Double-1>',lambda e:self.edit())
        footer=tk.Frame(self,bg='#185c37');footer.pack(fill='x',side='bottom')
        self.status=tk.StringVar(value='就绪 · 支持 PDF / 图片 / 整个文件夹，在本机处理。')
        tk.Label(footer,textvariable=self.status,bg='#185c37',fg='white',anchor='w',wraplength=1160,padx=24,pady=12).pack(fill='x')
        self.register_drop(self)
        self.after(200,self.poll)
    def register_drop(self,widget):
        widget.drop_target_register(DND_FILES);widget.dnd_bind('<<Drop>>',self.on_drop)
        widget.dnd_bind('<<DropEnter>>',lambda event:COPY)
        for child in widget.winfo_children():self.register_drop(child)
    def on_drop(self,event):
        try:self.import_paths(self.tk.splitlist(event.data))
        except Exception as e:self.status.set('无法导入拖入的内容：'+str(e))
        return COPY
    def save_amounts(self):
        p=filedialog.asksaveasfilename(title='导出逐票金额清单',defaultextension='.csv',filetypes=[('Excel可打开的金额清单','*.csv')])
        if p:
            try:write_amount_list(self.invoices,p);self.status.set('已导出逐票金额、已识别合计及金额待识别数量。')
            except Exception as e:messagebox.showerror('清单导出失败',str(e))
    def poll(self):
        try:
            while True:
                kind,data=self.q.get_nowait()
                if kind=='invoice':
                    if not any(i['hash']==data['hash'] or (data['number'] and i['number']==data['number']) for i in self.invoices):self.invoices.append(data);self.refresh()
                elif kind=='error': messagebox.showerror('处理提示',data)
                elif kind=='progress':self.status.set(data)
                elif kind=='done':self.importing=False;self.pending_paths.clear();self.status.set(data)
        except queue.Empty: pass
        self.after(200,self.poll)
    def refresh(self):
        self.tree.delete(*self.tree.get_children())
        for n,i in enumerate(self.invoices):
            try: amount=money(i['amount']); warn=amount>1000
            except Exception: amount='待核对'; warn=False
            try: validate(i); state='材料齐全，可分组'
            except Exception as e: state=str(e)
            if warn: state='超过1000元；'+state
            self.tree.insert('', 'end',iid=str(n),values=(Path(i['source']).name,i['number'] or '待识别',f'{amount:,.2f}' if isinstance(amount,Decimal) else amount,len(i['items']),f'{len(i["photos"])}张' if i.get('photos') else '不配图',state),tags=('warn',) if warn else ())
        self.stats=amount_summary(self.invoices)
        self.total_text.set(f'已识别金额合计  ¥ {self.stats["amount"]:,.2f}')
        self.count_text.set(f'共 {self.stats["count"]} 张  |  已识别金额 {self.stats["known"]} 张  |  金额待识别 {self.stats["unknown"]} 张  |  超1000元 {self.stats["alerts"]} 张')
    def selected(self):
        s=self.tree.selection(); return int(s[0]) if s else None
    def add(self):
        paths=filedialog.askopenfilenames(filetypes=[('发票','*.pdf *.png *.jpg *.jpeg *.bmp *.tif *.tiff')])
        if paths:self.import_paths(paths)
    def add_folder(self):
        path=filedialog.askdirectory(title='选择整个发票文件夹（包括子文件夹）')
        if path:self.import_paths([path])
    def import_paths(self,inputs):
        if self.importing:self.status.set('正在导入上一批，请完成后再拖入或选择下一批。');return
        self.importing=True;self.status.set('正在扫描文件及文件夹…')
        def worker():
            paths,skipped=collect_invoices(inputs);duplicates=0;failed=0;added=0
            hashes={i['hash'] for i in self.invoices};numbers={i['number'] for i in self.invoices if i['number']}
            for p in paths:
                self.q.put(('progress',f'正在识别 {added+duplicates+failed+1}/{len(paths)}：{Path(p).name}'))
                try:
                    digest=hashlib.sha256(Path(p).read_bytes()).hexdigest()
                    if digest in hashes:duplicates+=1;continue
                    try:data=extract(p)
                    except Exception as e:data=dict(source=p,hash=digest,number='',buyer='',amount='',category='采购',items=[],text=str(e),confirmed=False)
                    if data['number'] and data['number'] in numbers:duplicates+=1;continue
                    hashes.add(digest)
                    if data['number']:numbers.add(data['number'])
                    self.q.put(('invoice',data));added+=1
                except OSError:failed+=1
            self.q.put(('done',f'导入完成：新增 {added} 张，重复跳过 {duplicates} 张，不支持/无法访问 {len(skipped)+failed} 个。请核对识别金额与明细。'))
        threading.Thread(target=worker,daemon=True).start()
    def open_source(self):
        n=self.selected()
        if n is not None:
            try: os.startfile(self.invoices[n]['source'])
            except Exception as e: messagebox.showerror('无法打开',str(e))
    def remove(self):
        n=self.selected()
        if n is not None: self.invoices.pop(n); self.refresh()
    def edit(self):
        n=self.selected()
        if n is None:return
        inv=deepcopy(self.invoices[n]); win=tk.Toplevel(self); win.title('核对发票与资料'); win.geometry('1120x800'); win.grab_set()
        fields={}; top=ttk.Frame(win); top.pack(fill='x',padx=15,pady=10)
        if hasattr(self,'current_entry') and self.current_entry():ttk.Label(win,text='所属报销项：'+self.current_entry()['name'],foreground='#217346').pack(anchor='w',padx=20)
        for nfield,(k,label) in enumerate([('number','发票号码'),('buyer','购买方'),('amount','整票含税金额'),('category','费用/物品类别'),('seller','销售方名称'),('seller_tax','销售方税号'),('date','开票日期 YYYY-MM-DD'),('purpose','差旅费用说明（可选）')]):
            r,c=divmod(nfield,4);ttk.Label(top,text=label).grid(row=r*2,column=c,sticky='w',padx=6,pady=(6,0));v=tk.StringVar(value=inv.get(k,''));fields[k]=v;ttk.Entry(top,textvariable=v,width=27).grid(row=r*2+1,column=c,sticky='ew',padx=6)
        win.invoice_fields=fields
        ttk.Label(win,text='每行一个采购项目，小计填含税金额。每张发票可选一张或多张图片；直接下一张表示不配图，不留图片空档。').pack(anchor='w',padx=15)
        picbar=ttk.Frame(win);picbar.pack(fill='x',padx=15,pady=8)
        picinfo=tk.StringVar(value=f'已选 {len(inv.get("photos",[]))} 张图片' if inv.get('photos') else '不配图')
        def choosepic():
            p=filedialog.askopenfilenames(parent=win,filetypes=[('发票对应图片','*.jpg *.jpeg *.png *.bmp')])
            if p:inv['photos']=list(dict.fromkeys(inv.get('photos',[])+list(p)));picinfo.set(f'已选 {len(inv["photos"])} 张图片')
        def clearpic():inv['photos']=[];picinfo.set('不配图')
        ttk.Button(picbar,text='添加本票图片（可多选）',command=choosepic).pack(side='left')
        ttk.Button(picbar,text='不配图 / 清除图片',command=clearpic).pack(side='left',padx=8)
        ttk.Label(picbar,textvariable=picinfo).pack(side='left')
        if getattr(self,'kind','')=='差旅报销':picbar.pack_forget()
        tr=ttk.Treeview(win,columns=('name','model','qty','amount'),show='headings',height=8)
        for c,l,w in [('name','品名',300),('model','规格型号',180),('qty','数量',70),('amount','含税小计',100)]:tr.heading(c,text=l);tr.column(c,width=w)
        tr.pack(fill='x',padx=15,pady=10)
        def redraw():
            tr.delete(*tr.get_children())
            for k,x in enumerate(inv['items']):tr.insert('','end',iid=str(k),values=(x['name'],x['model'],x['qty'],x['amount']))
        def itemedit(new=False):
            s=tr.selection()
            if not new and not s:return
            idx=None if new else int(s[0]); item=dict(name='',model='',qty='',amount='',photos=[]) if new else deepcopy(inv['items'][idx])
            sub=tk.Toplevel(win); sub.title('采购项目'); sub.geometry('550x470'); sub.grab_set(); vs={}
            for k,l in [('name','物品名称'),('model','规格型号'),('qty','数量'),('amount','含税金额')]:
                ttk.Label(sub,text=l).pack(anchor='w',padx=20); v=tk.StringVar(value=item[k]); vs[k]=v;ttk.Entry(sub,textvariable=v,width=65).pack(padx=20,pady=3)
            def ok():
                item.update({k:v.get().strip() for k,v in vs.items()})
                try:
                    money(item['amount'])
                    if not item['name'] or Decimal(item['qty'])<=0:raise ValueError('请填写名称和正确数量')
                except Exception as e:messagebox.showerror('请核对',str(e),parent=sub);return
                if idx is None:inv['items'].append(item)
                else:inv['items'][idx]=item
                redraw();sub.destroy();win.grab_set()
            ttk.Button(sub,text='保存项目',command=ok).pack(pady=12)
        bar=ttk.Frame(win);bar.pack(fill='x',padx=15)
        def delete():
            s=tr.selection()
            if s:inv['items'].pop(int(s[0]));redraw()
        for l,c in [('新增项目',lambda:itemedit(True)),('编辑项目',itemedit),('删除项目',delete)]:ttk.Button(bar,text=l,command=c).pack(side='left',padx=4)
        ttk.Label(win,text='识别原文（供核对；无法可靠识别的字段留空，不自动编造）').pack(anchor='w',padx=15,pady=(12,0))
        txt=tk.Text(win,height=9,wrap='word');txt.insert('1.0',inv.get('text',''));txt.configure(state='disabled');txt.pack(fill='both',expand=True,padx=15,pady=5)
        confirm=tk.BooleanVar(value=inv.get('confirmed',False));ttk.Checkbutton(win,text='我已核对票号、金额、销售方、日期与当前报销类型；采购明细与整票金额一致，差旅可不录采购明细。',variable=confirm).pack(anchor='w',padx=15)
        win.confirmed_var=confirm
        def done(next_invoice=False):
            inv.update({k:v.get().strip() for k,v in fields.items()});inv['confirmed']=confirm.get()
            if any(k!=n and x['number'] and x['number']==inv['number'] for k,x in enumerate(self.invoices)):messagebox.showerror('重复发票','该票号已存在',parent=win);return
            self.invoices[n]=inv;self.refresh();win.destroy()
            next_index=self.next_review_index(n) if hasattr(self,'next_review_index') else (n+1 if n+1<len(self.invoices) else None)
            if next_invoice and next_index is not None:
                self.tree.selection_set(str(next_index));self._review_continuing=True;self.edit()
            elif hasattr(self,'_review_indices'):self._review_indices=[]
        bottom=ttk.Frame(win);bottom.pack(pady=12)
        ttk.Button(bottom,text='保存核对结果',command=done).pack(side='left',padx=8)
        ttk.Button(bottom,text='保存并下一张',command=lambda:done(True)).pack(side='left',padx=8);redraw()
        win.save_review=done
        return win
    def save(self):
        p=filedialog.asksaveasfilename(defaultextension='.json',filetypes=[('项目文件','*.json')])
        if p:
            Path(p).write_text(json.dumps(dict(version=1,invoices=self.invoices),ensure_ascii=False,indent=2),encoding='utf-8');self.status.set('项目已保存。项目记录引用原件路径，请勿移动发票和照片。')
    def load(self):
        p=filedialog.askopenfilename(filetypes=[('项目文件','*.json')])
        if p:
            try:
                data=json.loads(Path(p).read_text(encoding='utf-8'));assert data['version']==1 and isinstance(data['invoices'],list);self.invoices=data['invoices'];self.refresh()
            except Exception as e:messagebox.showerror('无法读取项目',str(e))
    def generate(self):
        try:batches,blocked=group(self.invoices)
        except Exception as e:messagebox.showerror('数据有误',str(e));return
        if not batches:messagebox.showinfo('尚不能导出','请先完成发票及明细核对。\n'+'\n'.join(f'{x.get("number") or Path(x["source"]).name}：{e}' for x,e in blocked));return
        lines=[f'第{k}份：{b["count"]}项，{b["amount"]:.2f}元，{len(b["invoices"])}张发票' for k,b in enumerate(batches,1)]
        if blocked:lines+=['\n以下发票暂不纳入：']+[f'{x.get("number") or Path(x["source"]).name}：{e}' for x,e in blocked]
        if not messagebox.askokcancel('分组预览','\n'.join(lines)+'\n\n导出以上材料齐全的分组？'):return
        p=filedialog.askdirectory(title='选择输出文件夹')
        if not p:return
        try:
            dest=export(batches,p);self.status.set('已导出 '+str(dest));os.startfile(dest);messagebox.showinfo('已生成','已生成 Word、发票和照片归档包。\n请在 Word 中预览核对分页和图片；验收结论、日期和签字保持空白。')
        except Exception as e:messagebox.showerror('导出未完成',str(e))
if __name__=='__main__':
    if len(sys.argv)==3 and sys.argv[1]=='--self-test':
        import regression
        try: regression.run(sys.argv[2])
        except Exception:
            import traceback
            Path(sys.argv[2]).mkdir(parents=True,exist_ok=True)
            (Path(sys.argv[2])/'error.txt').write_text(traceback.format_exc(),encoding='utf-8')
            raise
    elif len(sys.argv)==3 and sys.argv[1]=='--smoke-test':
        root=App();root.update();root.destroy()
        assert (BASE/'template.docx').is_file() and (BASE/'ocr.ps1').is_file()
        Path(sys.argv[2]).write_text('PASS: packaged UI, PDF engine, DOCX engine, template and OCR resources',encoding='utf-8')
    else: App().mainloop()
